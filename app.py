import os
import json
import datetime
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import customtkinter as ctk
import pandas as pd

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Configuración Global de Apariencia
ctk.set_appearance_mode("Light")
ctk.set_default_color_theme("blue")

CARPETA_DATOS = os.path.join(os.path.expanduser("~"), "ControlDisciplinarioEscritorio")
RUTA_BD = os.path.join(CARPETA_DATOS, "base_datos.json")

def inicializar_bd():
    if not os.path.exists(CARPETA_DATOS):
        os.makedirs(CARPETA_DATOS)
    if not os.path.exists(RUTA_BD):
        datos_iniciales = {"personal": {}, "boletas": []}
        with open(RUTA_BD, "w", encoding="utf-8") as f:
            json.dump(datos_iniciales, f, ensure_ascii=False, indent=4)

def cargar_bd():
    inicializar_bd()
    try:
        with open(RUTA_BD, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {"personal": {}, "boletas": []}

def guardar_bd(datos):
    inicializar_bd()
    with open(RUTA_BD, "r+", encoding="utf-8") as f:
        f.seek(0)
        json.dump(datos, f, ensure_ascii=False, indent=4)
        f.truncate()

# ==========================================
# --- MODAL: BUSCADOR GENERAL Y FALTAS ---
# ==========================================
class ModalBuscadorGeneral(ctk.CTkToplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.parent = parent
        self.title("🔍 Buscador General de Personal y Faltas Disciplinarias")
        self.geometry("1100x680")
        self.grab_set()

        self.COLOR_CARD = "#FFFFFF"
        self.COLOR_BORDER = "#CBD5E1"
        self.COLOR_TEXT_MAIN = "#0F172A"

        self.setup_ui()
        self.filtrar_resultados()

    def setup_ui(self):
        header_frame = ctk.CTkFrame(self, fg_color="#0F172A", corner_radius=0)
        header_frame.pack(fill="x", padx=0, pady=0)

        lbl_t = ctk.CTkLabel(header_frame, text="🔍 Buscador Global de Personal y Historial de Faltas", font=ctk.CTkFont(size=20, weight="bold"), text_color="#FFFFFF")
        lbl_t.pack(anchor="w", padx=20, pady=(15, 5))

        self.ent_search = ctk.CTkEntry(
            header_frame, 
            placeholder_text="Escriba Cédula, Nombre o cualquier texto para filtrar...", 
            height=42, 
            font=ctk.CTkFont(size=14)
        )
        self.ent_search.pack(fill="x", padx=20, pady=(5, 15))
        self.ent_search.bind("<KeyRelease>", self.filtrar_resultados)
        self.ent_search.focus_set()

        table_card = ctk.CTkFrame(self, corner_radius=12, fg_color=self.COLOR_CARD, border_width=1, border_color=self.COLOR_BORDER)
        table_card.pack(fill="both", expand=True, padx=20, pady=20)

        self.tv_modal = ttk.Treeview(table_card, columns=("Nro", "ID", "Nombre", "Tipo", "Fecha", "Gravedad", "Motivo"), show="headings")
        self.tv_modal.heading("Nro", text="N°")
        self.tv_modal.heading("ID", text="Cédula / ID")
        self.tv_modal.heading("Nombre", text="Nombre Completo")
        self.tv_modal.heading("Tipo", text="Tipo Boleta")
        self.tv_modal.heading("Fecha", text="Fecha")
        self.tv_modal.heading("Gravedad", text="Gravedad")
        self.tv_modal.heading("Motivo", text="Motivo / Descripción")

        self.tv_modal.column("Nro", width=50, anchor="center")
        self.tv_modal.column("ID", width=110, anchor="center")
        self.tv_modal.column("Nombre", width=220)
        self.tv_modal.column("Tipo", width=160)
        self.tv_modal.column("Fecha", width=100, anchor="center")
        self.tv_modal.column("Gravedad", width=110, anchor="center")
        self.tv_modal.column("Motivo", width=280)

        self.tv_modal.tag_configure("par", background="#F8FAFC")
        self.tv_modal.tag_configure("impar", background="#FFFFFF")

        sb = ttk.Scrollbar(table_card, orient="vertical", command=self.tv_modal.yview)
        self.tv_modal.configure(yscrollcommand=sb.set)

        self.tv_modal.pack(side="left", fill="both", expand=True, padx=(15, 0), pady=15)
        sb.pack(side="right", fill="y", padx=(0, 15), pady=15)

    def filtrar_resultados(self, event=None):
        for item in self.tv_modal.get_children():
            self.tv_modal.delete(item)

        q = self.ent_search.get().lower().strip()
        iconos = {"Baja": "🟢 Baja", "Media": "🟡 Media", "Alta": "🟠 Alta", "Urgente": "🔴 Urgente"}

        idx = 0
        for b in reversed(self.parent.bd["boletas"]):
            pid = str(b.get("persona_id", ""))
            pdata = self.parent.bd["personal"].get(pid, {})
            pnombre = str(pdata.get("nombre", ""))
            tipo = str(b.get("tipo", ""))
            motivo = str(b.get("motivo", ""))

            if not q or (q in pid.lower() or q in pnombre.lower() or q in tipo.lower() or q in motivo.lower()):
                tag = "par" if idx % 2 == 0 else "impar"
                grav_txt = iconos.get(b.get("gravedad"), b.get("gravedad"))
                self.tv_modal.insert("", "end", values=(
                    b.get("id_boleta"), pid, pnombre, tipo, b.get("fecha"), grav_txt, motivo
                ), tags=(tag,))
                idx += 1


# ==========================================
# --- MODAL: CONTEO Y MÉTRICAS POR USUARIO ---
# ==========================================
# ==========================================
# --- MODAL: CONTEO Y MÉTRICAS POR USUARIO ---
# ==========================================
# ==========================================
# --- MODAL: CONTEO Y MÉTRICAS POR USUARIO ---
# ==========================================
class ModalTotalBoletasUsuario(ctk.CTkToplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.parent = parent
        self.title("📊 Total de Boletas y Métricas por Personal")
        self.geometry("980x640")
        self.grab_set()

        self.COLOR_CARD = "#FFFFFF"
        self.COLOR_BORDER = "#CBD5E1"

        self.setup_ui()
        self.actualizar_combo_usuarios()

    def setup_ui(self):
        top_bar = ctk.CTkFrame(self, fg_color="#1E293B", corner_radius=0)
        top_bar.pack(fill="x", padx=0, pady=0)

        ctk.CTkLabel(top_bar, text="Filtrar Personal:", font=ctk.CTkFont(size=14, weight="bold"), text_color="#FFFFFF").pack(side="left", padx=(20, 5), pady=15)
        
        self.ent_search_user = ctk.CTkEntry(top_bar, placeholder_text="Buscar Cédula o Nombre...", width=200, height=38)
        self.ent_search_user.pack(side="left", padx=5, pady=15)

        self.combo_usuarios = ctk.CTkComboBox(top_bar, width=320, height=38, state="readonly", command=self.calcular_metricas_usuario)
        self.combo_usuarios.pack(side="left", padx=10, pady=15)

        # Se asigna el binding después de crear self.combo_usuarios para evitar el AttributeError
        self.ent_search_user.bind("<KeyRelease>", self.actualizar_combo_usuarios)

        self.kpi_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.kpi_frame.pack(fill="x", padx=20, pady=15)
        self.kpi_frame.columnconfigure((0, 1, 2), weight=1)

        self.lbl_tot_boletas = self.crear_card_stat(self.kpi_frame, 0, "Total Boletas", "0", "#2563EB", "#EFF6FF")
        self.lbl_tot_urgentes = self.crear_card_stat(self.kpi_frame, 1, "Boletas Altas/Urgentes", "0", "#DC2626", "#FEF2F2")
        self.lbl_tot_tipo = self.crear_card_stat(self.kpi_frame, 2, "Tipo Frecuente", "-", "#0D9488", "#CCFBF1")

        table_card = ctk.CTkFrame(self, corner_radius=12, fg_color=self.COLOR_CARD, border_width=1, border_color=self.COLOR_BORDER)
        table_card.pack(fill="both", expand=True, padx=20, pady=(0, 20))

        self.tv_user_boletas = ttk.Treeview(table_card, columns=("Nro", "Tipo", "Fecha", "Gravedad", "Motivo"), show="headings")
        self.tv_user_boletas.heading("Nro", text="N°")
        self.tv_user_boletas.heading("Tipo", text="Tipo Boleta")
        self.tv_user_boletas.heading("Fecha", text="Fecha")
        self.tv_user_boletas.heading("Gravedad", text="Gravedad")
        self.tv_user_boletas.heading("Motivo", text="Motivo / Descripción")

        self.tv_user_boletas.column("Nro", width=60, anchor="center")
        self.tv_user_boletas.column("Tipo", width=180)
        self.tv_user_boletas.column("Fecha", width=110, anchor="center")
        self.tv_user_boletas.column("Gravedad", width=120, anchor="center")
        self.tv_user_boletas.column("Motivo", width=340)

        self.tv_user_boletas.pack(fill="both", expand=True, padx=15, pady=15)

    def crear_card_stat(self, parent, col, titulo, valor, color_txt, color_bg):
        card = ctk.CTkFrame(parent, corner_radius=10, fg_color=color_bg, border_width=1, border_color=self.COLOR_BORDER)
        card.grid(row=0, column=col, padx=8, pady=5, sticky="ew")

        lbl_v = ctk.CTkLabel(card, text=valor, font=ctk.CTkFont(size=24, weight="bold"), text_color=color_txt)
        lbl_v.pack(padx=15, pady=(12, 0))

        lbl_t = ctk.CTkLabel(card, text=titulo, font=ctk.CTkFont(size=12, weight="bold"), text_color="#64748B")
        lbl_t.pack(padx=15, pady=(2, 12))
        return lbl_v

    def actualizar_combo_usuarios(self, event=None):
        if not hasattr(self, 'combo_usuarios'):
            return

        q = self.ent_search_user.get().lower().strip() if hasattr(self, 'ent_search_user') else ""
        lista = []
        for pid, pdata in self.parent.bd["personal"].items():
            nombre = pdata.get('nombre', '')
            if not q or q in str(pid).lower() or q in nombre.lower():
                lista.append(f"{pid} - {nombre}")

        if lista:
            self.combo_usuarios.configure(values=lista)
            self.combo_usuarios.set(lista[0])
            self.calcular_metricas_usuario(lista[0])
        else:
            self.combo_usuarios.configure(values=["No hay coincidencias"])
            self.combo_usuarios.set("No hay coincidencias")
            self.calcular_metricas_usuario("")

    def calcular_metricas_usuario(self, choice):
        for item in self.tv_user_boletas.get_children():
            self.tv_user_boletas.delete(item)

        if " - " not in choice:
            self.lbl_tot_boletas.configure(text="0")
            self.lbl_tot_urgentes.configure(text="0")
            self.lbl_tot_tipo.configure(text="-")
            return

        pid = choice.split(" - ")[0].strip()
        boletas_usr = [b for b in self.parent.bd["boletas"] if str(b.get("persona_id")) == pid]

        tot = len(boletas_usr)
        criticas = sum(1 for b in boletas_usr if b.get("gravedad") in ["Alta", "Urgente"])

        tipos = [b.get("tipo") for b in boletas_usr if b.get("tipo")]
        frecuente = max(set(tipos), key=tipos.count) if tipos else "-"

        self.lbl_tot_boletas.configure(text=str(tot))
        self.lbl_tot_urgentes.configure(text=str(criticas))
        self.lbl_tot_tipo.configure(text=frecuente)

        iconos = {"Baja": "🟢 Baja", "Media": "🟡 Media", "Alta": "🟠 Alta", "Urgente": "🔴 Urgente"}
        for b in reversed(boletas_usr):
            grav_txt = iconos.get(b.get("gravedad"), b.get("gravedad"))
            self.tv_user_boletas.insert("", "end", values=(
                b.get("id_boleta"), b.get("tipo"), b.get("fecha"), grav_txt, b.get("motivo")
            ))
    def __init__(self, parent):
        super().__init__(parent)
        self.parent = parent
        self.title("📊 Total de Boletas y Métricas por Personal")
        self.geometry("980x640")
        self.grab_set()

        self.COLOR_CARD = "#FFFFFF"
        self.COLOR_BORDER = "#CBD5E1"

        self.setup_ui()
        self.actualizar_combo_usuarios()

    def setup_ui(self):
        top_bar = ctk.CTkFrame(self, fg_color="#1E293B", corner_radius=0)
        top_bar.pack(fill="x", padx=0, pady=0)

        ctk.CTkLabel(top_bar, text="Filtrar Personal:", font=ctk.CTkFont(size=14, weight="bold"), text_color="#FFFFFF").pack(side="left", padx=(20, 5), pady=15)
        
        # Buscador directo para el combobox
        self.ent_search_user = ctk.CTkEntry(top_bar, placeholder_text="Buscar Cédula o Nombre...", width=200, height=38)
        self.ent_search_user.pack(side="left", padx=5, pady=15)
        self.ent_search_user.bind("<KeyRelease>", self.actualizar_combo_usuarios)

        self.combo_usuarios = ctk.CTkComboBox(top_bar, width=320, height=38, state="readonly", command=self.calcular_metricas_usuario)
        self.combo_usuarios.pack(side="left", padx=10, pady=15)

        self.kpi_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.kpi_frame.pack(fill="x", padx=20, pady=15)
        self.kpi_frame.columnconfigure((0, 1, 2), weight=1)

        self.lbl_tot_boletas = self.crear_card_stat(self.kpi_frame, 0, "Total Boletas", "0", "#2563EB", "#EFF6FF")
        self.lbl_tot_urgentes = self.crear_card_stat(self.kpi_frame, 1, "Boletas Altas/Urgentes", "0", "#DC2626", "#FEF2F2")
        self.lbl_tot_tipo = self.crear_card_stat(self.kpi_frame, 2, "Tipo Frecuente", "-", "#0D9488", "#CCFBF1")

        table_card = ctk.CTkFrame(self, corner_radius=12, fg_color=self.COLOR_CARD, border_width=1, border_color=self.COLOR_BORDER)
        table_card.pack(fill="both", expand=True, padx=20, pady=(0, 20))

        self.tv_user_boletas = ttk.Treeview(table_card, columns=("Nro", "Tipo", "Fecha", "Gravedad", "Motivo"), show="headings")
        self.tv_user_boletas.heading("Nro", text="N°")
        self.tv_user_boletas.heading("Tipo", text="Tipo Boleta")
        self.tv_user_boletas.heading("Fecha", text="Fecha")
        self.tv_user_boletas.heading("Gravedad", text="Gravedad")
        self.tv_user_boletas.heading("Motivo", text="Motivo / Descripción")

        self.tv_user_boletas.column("Nro", width=60, anchor="center")
        self.tv_user_boletas.column("Tipo", width=180)
        self.tv_user_boletas.column("Fecha", width=110, anchor="center")
        self.tv_user_boletas.column("Gravedad", width=120, anchor="center")
        self.tv_user_boletas.column("Motivo", width=340)

        self.tv_user_boletas.pack(fill="both", expand=True, padx=15, pady=15)

    def actualizar_combo_usuarios(self, event=None):
        q = self.ent_search_user.get().lower().strip() if hasattr(self, 'ent_search_user') else ""
        lista = []
        for pid, pdata in self.parent.bd["personal"].items():
            nombre = pdata.get('nombre', '')
            if not q or q in str(pid).lower() or q in nombre.lower():
                lista.append(f"{pid} - {nombre}")

        if lista:
            self.combo_usuarios.configure(values=lista)
            self.combo_usuarios.set(lista[0])
            self.calcular_metricas_usuario(lista[0])
        else:
            self.combo_usuarios.configure(values=["No hay coincidencias"])
            self.combo_usuarios.set("No hay coincidencias")
            self.calcular_metricas_usuario("")

# ==========================================
# --- PANEL PRINCIPAL DASHBOARD ---
# ==========================================
class DashboardAdmin(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Sistema de Control Disciplinario — Panel de Control")

        # Maximizar ventana sin ocultar barra de tareas ni botones de control
        try:
            self.state("zoomed")  # Funciona en Windows
        except Exception:
            self.attributes("-zoomed", True)  # Funciona en Linux (Parrot OS / X11)

        self.minsize(1100, 700)

        self.COLOR_BG = "#F1F5F9"
        self.COLOR_SIDEBAR = "#0F172A"
        self.COLOR_CARD = "#FFFFFF"
        self.COLOR_BORDER = "#CBD5E1"
        self.COLOR_TEXT_MAIN = "#0F172A"
        self.COLOR_TEXT_MUTED = "#64748B"

        self.configure(fg_color=self.COLOR_BG)

        self.bd = cargar_bd()
        self.boleta_id_actual = None
        self.persona_seleccionada_boleta = None

        self.estilar_tablas()

        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self.setup_sidebar()
        self.setup_contenido_principal()
        self.mostrar_vista("dashboard")

    def estilar_tablas(self):
        style = ttk.Style()
        style.theme_use("clam")
        
        style.configure("Treeview.Heading",
                        font=("Segoe UI", 10, "bold"),
                        background="#E2E8F0",
                        foreground="#0F172A",
                        relief="flat",
                        padding=10)

        style.configure("Treeview",
                        font=("Segoe UI", 10),
                        background="#FFFFFF",
                        foreground="#0F172A",
                        rowheight=36,
                        fieldbackground="#FFFFFF",
                        borderwidth=0)

        style.map("Treeview",
                  background=[("selected", "#2563EB")],
                  foreground=[("selected", "#FFFFFF")])

    def setup_sidebar(self):
        self.sidebar = ctk.CTkFrame(self, width=260, corner_radius=0, fg_color=self.COLOR_SIDEBAR)
        self.sidebar.grid(row=0, column=0, sticky="nsew")
        self.sidebar.grid_rowconfigure(5, weight=1)

        lbl_logo = ctk.CTkLabel(
            self.sidebar, 
            text="🛡️ AdminControl", 
            font=ctk.CTkFont(size=22, weight="bold"),
            text_color="#FFFFFF"
        )
        lbl_logo.grid(row=0, column=0, padx=20, pady=(30, 35), sticky="w")

        self.btn_nav_dash = self.crear_btn_nav("📊 Dashboard", "dashboard", 1)
        self.btn_nav_pers = self.crear_btn_nav("👤 Personal", "personal", 2)
        self.btn_nav_boletas = self.crear_btn_nav("📝 Emitir Boleta", "boletas", 3)
        self.btn_nav_reportes = self.crear_btn_nav("📋 Historial General", "reportes", 4)

        btn_excel_side = ctk.CTkButton(
            self.sidebar, text="📊 Exportar a Excel", fg_color="#10B981", hover_color="#059669",
            text_color="white", font=ctk.CTkFont(size=14, weight="bold"), height=42,
            command=self.exportar_excel
        )
        btn_excel_side.grid(row=6, column=0, padx=16, pady=25, sticky="ew")

    def crear_btn_nav(self, text_label, vista, row_idx):
        btn = ctk.CTkButton(
            self.sidebar, text=text_label, anchor="w", fg_color="transparent", 
            text_color="#94A3B8", hover_color="#1E293B", font=ctk.CTkFont(size=14, weight="bold"),
            height=44, command=lambda: self.mostrar_vista(vista)
        )
        btn.grid(row=row_idx, column=0, padx=14, pady=4, sticky="ew")
        return btn

    def setup_contenido_principal(self):
        self.main_container = ctk.CTkFrame(self, corner_radius=0, fg_color="transparent")
        self.main_container.grid(row=0, column=1, sticky="nsew", padx=30, pady=25)

        self.vista_dash = ctk.CTkFrame(self.main_container, fg_color="transparent")
        self.vista_pers = ctk.CTkFrame(self.main_container, fg_color="transparent")
        self.vista_boletas = ctk.CTkFrame(self.main_container, fg_color="transparent")
        self.vista_reportes = ctk.CTkFrame(self.main_container, fg_color="transparent")

        self.setup_vista_dashboard()
        self.setup_vista_personal()
        self.setup_vista_boletas()
        self.setup_vista_reportes()

    def mostrar_vista(self, nombre_vista):
        for v in [self.vista_dash, self.vista_pers, self.vista_boletas, self.vista_reportes]:
            v.pack_forget()

        btns = {
            "dashboard": self.btn_nav_dash,
            "personal": self.btn_nav_pers,
            "boletas": self.btn_nav_boletas,
            "reportes": self.btn_nav_reportes
        }
        for k, btn in btns.items():
            if k == nombre_vista:
                btn.configure(fg_color="#1E293B", text_color="#FFFFFF")
            else:
                btn.configure(fg_color="transparent", text_color="#94A3B8")

        if nombre_vista == "dashboard":
            self.actualizar_kpis()
            self.actualizar_tabla_resumen()
            self.vista_dash.pack(fill="both", expand=True)
        elif nombre_vista == "personal":
            self.vista_pers.pack(fill="both", expand=True)
        elif nombre_vista == "boletas":
            self.filtrar_combo_personal_boleta()
            self.vista_boletas.pack(fill="both", expand=True)
        elif nombre_vista == "reportes":
            self.actualizar_combo_filtro_usuarios()
            self.actualizar_tabla_reportes()
            self.vista_reportes.pack(fill="both", expand=True)

    # ==========================================
    # --- VISTA 1: DASHBOARD CON MODALES ---
    # ==========================================
    def setup_vista_dashboard(self):
        top_frame = ctk.CTkFrame(self.vista_dash, fg_color="transparent")
        top_frame.pack(fill="x", pady=(0, 15))

        lbl_t = ctk.CTkLabel(top_frame, text="Resumen de Actividad General", font=ctk.CTkFont(size=24, weight="bold"), text_color=self.COLOR_TEXT_MAIN)
        lbl_t.pack(side="left")

        action_card = ctk.CTkFrame(self.vista_dash, corner_radius=12, fg_color=self.COLOR_CARD, border_width=1, border_color=self.COLOR_BORDER)
        action_card.pack(fill="x", pady=(0, 20), padx=2)

        lbl_s = ctk.CTkLabel(action_card, text="🔍 Búsqueda e Historial de Personal:", font=ctk.CTkFont(size=14, weight="bold"), text_color=self.COLOR_TEXT_MAIN)
        lbl_s.pack(side="left", padx=15, pady=15)

        btn_modal_search = ctk.CTkButton(
            action_card, 
            text="🔎 Abrir Buscador en Modal", 
            fg_color="#2563EB", hover_color="#1D4ED8", 
            font=ctk.CTkFont(size=13, weight="bold"), 
            height=38,
            command=self.abrir_modal_buscador
        )
        btn_modal_search.pack(side="left", padx=10, pady=15)

        btn_modal_user_stats = ctk.CTkButton(
            action_card, 
            text="📊 Total Boletas por Usuario (Modal)", 
            fg_color="#0D9488", hover_color="#0F766E", 
            font=ctk.CTkFont(size=13, weight="bold"), 
            height=38,
            command=self.abrir_modal_total_usuario
        )
        btn_modal_user_stats.pack(side="left", padx=10, pady=15)

        kpi_frame = ctk.CTkFrame(self.vista_dash, fg_color="transparent")
        kpi_frame.pack(fill="x", pady=(0, 20))
        kpi_frame.columnconfigure((0, 1, 2, 3), weight=1)

        self.card1 = self.crear_kpi_card(kpi_frame, 0, "Total Personal", "0", "#2563EB", "#EFF6FF")
        self.card2 = self.crear_kpi_card(kpi_frame, 1, "Boletas Emitidas", "0", "#0D9488", "#CCFBF1")
        self.card3 = self.crear_kpi_card(kpi_frame, 2, "Casos Críticos", "0", "#E11D48", "#FFE4E6")
        self.card4 = self.crear_kpi_card(kpi_frame, 3, "Atenciones Mes", "0", "#D97706", "#FEF3C7")

        lbl_rec = ctk.CTkLabel(self.vista_dash, text="Últimas Boletas Emitidas", font=ctk.CTkFont(size=16, weight="bold"), text_color=self.COLOR_TEXT_MAIN)
        lbl_rec.pack(anchor="w", pady=(5, 10))

        card_table = ctk.CTkFrame(self.vista_dash, corner_radius=12, fg_color=self.COLOR_CARD, border_width=1, border_color=self.COLOR_BORDER)
        card_table.pack(fill="both", expand=True)

        self.tv_dash = ttk.Treeview(card_table, columns=("Nro", "Persona", "Tipo", "Fecha", "Gravedad"), show="headings", height=8)
        self.tv_dash.heading("Nro", text="N°")
        self.tv_dash.heading("Persona", text="Persona")
        self.tv_dash.heading("Tipo", text="Tipo de Boleta")
        self.tv_dash.heading("Fecha", text="Fecha")
        self.tv_dash.heading("Gravedad", text="Gravedad")

        self.tv_dash.column("Nro", width=60, anchor="center")
        self.tv_dash.column("Persona", width=280)
        self.tv_dash.column("Tipo", width=200)
        self.tv_dash.column("Fecha", width=120, anchor="center")
        self.tv_dash.column("Gravedad", width=130, anchor="center")

        self.tv_dash.tag_configure("par", background="#F8FAFC")
        self.tv_dash.tag_configure("impar", background="#FFFFFF")

        self.tv_dash.pack(fill="both", expand=True, padx=15, pady=15)

    def abrir_modal_buscador(self):
        ModalBuscadorGeneral(self)

    def abrir_modal_total_usuario(self):
        ModalTotalBoletasUsuario(self)

    def crear_kpi_card(self, parent, col, titulo, valor, color_texto, color_bg):
        card = ctk.CTkFrame(parent, corner_radius=12, fg_color=color_bg, border_width=1, border_color=self.COLOR_BORDER)
        card.grid(row=0, column=col, padx=8, pady=5, sticky="ew")
        
        lbl_val = ctk.CTkLabel(card, text=valor, font=ctk.CTkFont(size=32, weight="bold"), text_color=color_texto)
        lbl_val.pack(padx=18, pady=(18, 0))

        lbl_tit = ctk.CTkLabel(card, text=titulo, font=ctk.CTkFont(size=12, weight="bold"), text_color=self.COLOR_TEXT_MUTED)
        lbl_tit.pack(padx=18, pady=(2, 18))
        return lbl_val

    def actualizar_kpis(self):
        tot_pers = len(self.bd["personal"])
        tot_boletas = len(self.bd["boletas"])
        urgentes = sum(1 for b in self.bd["boletas"] if b.get("gravedad") in ["Alta", "Urgente"])
        mes_act = datetime.datetime.now().strftime("%m/%Y")
        del_mes = sum(1 for b in self.bd["boletas"] if b.get("fecha", "").endswith(mes_act))

        self.card1.configure(text=str(tot_pers))
        self.card2.configure(text=str(tot_boletas))
        self.card3.configure(text=str(urgentes))
        self.card4.configure(text=str(del_mes))

    def actualizar_tabla_resumen(self):
        for item in self.tv_dash.get_children():
            self.tv_dash.delete(item)

        datos = list(reversed(self.bd["boletas"][-8:]))
        for idx, b in enumerate(datos):
            pid = str(b.get("persona_id", ""))
            pnombre = self.bd["personal"].get(pid, {}).get("nombre", "Desconocido")
            tag = "par" if idx % 2 == 0 else "impar"
            
            grav = b.get("gravedad", "")
            iconos = {"Baja": "🟢 Baja", "Media": "🟡 Media", "Alta": "🟠 Alta", "Urgente": "🔴 Urgente"}
            grav_txt = iconos.get(grav, grav)

            self.tv_dash.insert("", "end", values=(b.get("id_boleta"), f"{pnombre} ({pid})", b.get("tipo"), b.get("fecha"), grav_txt), tags=(tag,))

    # ==========================================
    # --- VISTA 2: GESTIÓN DE PERSONAL ---
    # ==========================================
    def setup_vista_personal(self):
        lbl_t = ctk.CTkLabel(self.vista_pers, text="Administración de Personal", font=ctk.CTkFont(size=24, weight="bold"), text_color=self.COLOR_TEXT_MAIN)
        lbl_t.pack(anchor="w", pady=(0, 15))

        form_card = ctk.CTkFrame(self.vista_pers, corner_radius=12, fg_color=self.COLOR_CARD, border_width=1, border_color=self.COLOR_BORDER)
        form_card.pack(fill="x", pady=(0, 20))

        ctk.CTkLabel(form_card, text="Identificación / RUT:", text_color=self.COLOR_TEXT_MAIN, font=ctk.CTkFont(weight="bold")).grid(row=0, column=0, padx=(20, 5), pady=15, sticky="w")
        self.ent_pers_id = ctk.CTkEntry(form_card, placeholder_text="Ej. 12345678-K", width=220, height=36)
        self.ent_pers_id.grid(row=0, column=1, padx=5, pady=15, sticky="w")

        ctk.CTkLabel(form_card, text="Nombre Completo:", text_color=self.COLOR_TEXT_MAIN, font=ctk.CTkFont(weight="bold")).grid(row=0, column=2, padx=(20, 5), pady=15, sticky="w")
        self.ent_pers_nombre = ctk.CTkEntry(form_card, placeholder_text="Ej. Juan Pérez", width=260, height=36)
        self.ent_pers_nombre.grid(row=0, column=3, padx=5, pady=15, sticky="w")

        ctk.CTkLabel(form_card, text="Cargo / Curso:", text_color=self.COLOR_TEXT_MAIN, font=ctk.CTkFont(weight="bold")).grid(row=1, column=0, padx=(20, 5), pady=15, sticky="w")
        self.ent_pers_cargo = ctk.CTkEntry(form_card, placeholder_text="Ej. Operador / 4to Medio", width=220, height=36)
        self.ent_pers_cargo.grid(row=1, column=1, padx=5, pady=15, sticky="w")

        ctk.CTkLabel(form_card, text="Departamento:", text_color=self.COLOR_TEXT_MAIN, font=ctk.CTkFont(weight="bold")).grid(row=1, column=2, padx=(20, 5), pady=15, sticky="w")
        self.ent_pers_depto = ctk.CTkEntry(form_card, placeholder_text="Ej. Logística / Básico", width=260, height=36)
        self.ent_pers_depto.grid(row=1, column=3, padx=5, pady=15, sticky="w")

        btn_box = ctk.CTkFrame(form_card, fg_color="transparent")
        btn_box.grid(row=2, column=0, columnspan=4, pady=18)

        ctk.CTkButton(btn_box, text="💾 Guardar", fg_color="#2563EB", hover_color="#1D4ED8", width=130, height=36, command=self.guardar_personal).pack(side="left", padx=8)
        ctk.CTkButton(btn_box, text="🧹 Limpiar", fg_color="#64748B", hover_color="#475569", width=120, height=36, command=self.limpiar_form_personal).pack(side="left", padx=8)
        ctk.CTkButton(btn_box, text="🗑️ Eliminar", fg_color="#EF4444", hover_color="#DC2626", width=130, height=36, command=self.eliminar_personal).pack(side="left", padx=8)
        ctk.CTkButton(btn_box, text="📥 Importar Excel", fg_color="#0D9488", hover_color="#0F766E", width=150, height=36, command=self.importar_personal_excel).pack(side="left", padx=8)

        table_card = ctk.CTkFrame(self.vista_pers, corner_radius=12, fg_color=self.COLOR_CARD, border_width=1, border_color=self.COLOR_BORDER)
        table_card.pack(fill="both", expand=True)

        self.tv_personal = ttk.Treeview(table_card, columns=("ID", "Nombre", "Cargo", "Depto"), show="headings")
        self.tv_personal.heading("ID", text="ID / Cédula")
        self.tv_personal.heading("Nombre", text="Nombre Completo")
        self.tv_personal.heading("Cargo", text="Cargo / Curso")
        self.tv_personal.heading("Depto", text="Departamento / Sección")

        self.tv_personal.column("ID", width=150, anchor="center")
        self.tv_personal.column("Nombre", width=280)
        self.tv_personal.column("Cargo", width=200)
        self.tv_personal.column("Depto", width=200)

        self.tv_personal.tag_configure("par", background="#F8FAFC")
        self.tv_personal.tag_configure("impar", background="#FFFFFF")

        sb = ttk.Scrollbar(table_card, orient="vertical", command=self.tv_personal.yview)
        self.tv_personal.configure(yscrollcommand=sb.set)

        self.tv_personal.pack(side="left", fill="both", expand=True, padx=(15, 0), pady=15)
        sb.pack(side="right", fill="y", padx=(0, 15), pady=15)

        self.tv_personal.bind("<<TreeviewSelect>>", self.seleccionar_personal)
        self.actualizar_tabla_personal()

    def guardar_personal(self):
        pid = self.ent_pers_id.get().strip()
        nombre = self.ent_pers_nombre.get().strip()
        cargo = self.ent_pers_cargo.get().strip()
        depto = self.ent_pers_depto.get().strip()

        if not pid or not nombre:
            messagebox.showwarning("Atención", "Debe ingresar la Identificación y el Nombre.")
            return

        self.bd["personal"][pid] = {"nombre": nombre, "cargo": cargo, "depto": depto}
        guardar_bd(self.bd)
        messagebox.showinfo("Éxito", f"Personal '{nombre}' guardado.")
        self.limpiar_form_personal()
        self.actualizar_tabla_personal()

    def limpiar_form_personal(self):
        self.ent_pers_id.configure(state="normal")
        self.ent_pers_id.delete(0, "end")
        self.ent_pers_nombre.delete(0, "end")
        self.ent_pers_cargo.delete(0, "end")
        self.ent_pers_depto.delete(0, "end")

    def seleccionar_personal(self, event):
        sel = self.tv_personal.selection()
        if sel:
            vals = self.tv_personal.item(sel[0])["values"]
            self.limpiar_form_personal()
            self.ent_pers_id.insert(0, str(vals[0]))
            self.ent_pers_id.configure(state="disabled")
            self.ent_pers_nombre.insert(0, str(vals[1]))
            self.ent_pers_cargo.insert(0, str(vals[2]))
            self.ent_pers_depto.insert(0, str(vals[3]))

    def eliminar_personal(self):
        sel = self.tv_personal.selection()
        if not sel:
            messagebox.showwarning("Atención", "Seleccione un registro para eliminar.")
            return
        pid = str(self.tv_personal.item(sel[0])["values"][0])
        if messagebox.askyesno("Confirmar", f"¿Eliminar a la persona ID {pid}?"):
            if pid in self.bd["personal"]:
                del self.bd["personal"][pid]
                guardar_bd(self.bd)
                self.actualizar_tabla_personal()
                self.limpiar_form_personal()

    def actualizar_tabla_personal(self):
        for item in self.tv_personal.get_children():
            self.tv_personal.delete(item)
        for idx, (pid, pdata) in enumerate(self.bd["personal"].items()):
            tag = "par" if idx % 2 == 0 else "impar"
            self.tv_personal.insert("", "end", values=(pid, pdata.get("nombre", ""), pdata.get("cargo", ""), pdata.get("depto", "")), tags=(tag,))

    def importar_personal_excel(self):
        ruta = filedialog.askopenfilename(
            title="Seleccionar archivo Excel de Personal",
            filetypes=[("Archivos de Excel", "*.xlsx *.xls")]
        )
        if not ruta:
            return

        try:
            df = pd.read_excel(ruta)
            columnas = [str(c).lower().strip() for c in df.columns]

            def buscar_col(posibles):
                for p in posibles:
                    for idx, c in enumerate(columnas):
                        if p in c:
                            return df.columns[idx]
                return None

            col_id = buscar_col(["id", "rut", "cedula", "identificacion", "codigo"])
            col_nom = buscar_col(["nombre", "persona", "empleado", "trabajador"])
            col_car = buscar_col(["cargo", "curso", "puesto", "rol"])
            col_dep = buscar_col(["depto", "departamento", "seccion", "area"])

            if not col_id or not col_nom:
                messagebox.showerror("Error de Formato", "El archivo Excel debe contener al menos las columnas 'ID/Cédula' y 'Nombre'.")
                return

            registrados = 0
            for _, row in df.iterrows():
                pid = str(row[col_id]).strip()
                nombre = str(row[col_nom]).strip()
                cargo = str(row[col_car]).strip() if col_car and pd.notna(row[col_car]) else ""
                depto = str(row[col_dep]).strip() if col_dep and pd.notna(row[col_dep]) else ""

                if pid and nombre and pid.lower() != "nan" and nombre.lower() != "nan":
                    self.bd["personal"][pid] = {"nombre": nombre, "cargo": cargo, "depto": depto}
                    registrados += 1

            guardar_bd(self.bd)
            self.actualizar_tabla_personal()
            messagebox.showinfo("Importación Exitosa", f"Se importaron {registrados} registros de personal correctamente.")
        except Exception as e:
            messagebox.showerror("Error de Importación", f"No se pudo procesar el archivo Excel.\nDetalle: {str(e)}")

    # ==========================================
    # --- VISTA 3: EMISIÓN DE BOLETAS ---
    # ==========================================
    def setup_vista_boletas(self):
        lbl_t = ctk.CTkLabel(self.vista_boletas, text="Emisión de Boletas Disciplinarias", font=ctk.CTkFont(size=24, weight="bold"), text_color=self.COLOR_TEXT_MAIN)
        lbl_t.pack(anchor="w", pady=(0, 15))

        form_card = ctk.CTkFrame(self.vista_boletas, corner_radius=12, fg_color=self.COLOR_CARD, border_width=1, border_color=self.COLOR_BORDER)
        form_card.pack(fill="both", expand=True)

        ctk.CTkLabel(form_card, text="Buscar Personal:", text_color=self.COLOR_TEXT_MAIN, font=ctk.CTkFont(weight="bold")).grid(row=0, column=0, padx=(20, 5), pady=15, sticky="w")
        
        search_box = ctk.CTkFrame(form_card, fg_color="transparent")
        search_box.grid(row=0, column=1, padx=5, pady=15, sticky="w")

        self.ent_buscar_persona_boleta = ctk.CTkEntry(search_box, placeholder_text="Cédula, Nombre...", width=220, height=36)
        self.ent_buscar_persona_boleta.pack(side="left", padx=(0, 5))
        self.ent_buscar_persona_boleta.bind("<KeyRelease>", self.filtrar_combo_personal_boleta)

        self.combo_personas = ctk.CTkComboBox(search_box, width=280, height=36, state="readonly", command=self.al_seleccionar_persona_boleta)
        self.combo_personas.pack(side="left")

        self.lbl_info_persona = ctk.CTkLabel(
            form_card, 
            text="Cargo: - | Depto: -", 
            font=ctk.CTkFont(size=12, slant="italic"), 
            text_color=self.COLOR_TEXT_MUTED
        )
        self.lbl_info_persona.grid(row=0, column=2, columnspan=2, padx=15, pady=15, sticky="w")

        ctk.CTkLabel(form_card, text="Tipo de Boleta:", text_color=self.COLOR_TEXT_MAIN, font=ctk.CTkFont(weight="bold")).grid(row=1, column=0, padx=(20, 5), pady=15, sticky="w")
        self.combo_tipo_boleta = ctk.CTkComboBox(
            form_card, 
            values=["Aviso Disciplinario", "Llamado de Atención", "Infracción Leve", "Infracción Grave", "Citación Apoderado/Jefe", "Pase de Salida"], 
            width=300, height=36, state="readonly"
        )
        self.combo_tipo_boleta.grid(row=1, column=1, padx=5, pady=15, sticky="w")

        ctk.CTkLabel(form_card, text="Fecha de Emisión:", text_color=self.COLOR_TEXT_MAIN, font=ctk.CTkFont(weight="bold")).grid(row=1, column=2, padx=(20, 5), pady=15, sticky="w")
        self.ent_fecha = ctk.CTkEntry(form_card, width=220, height=36)
        self.ent_fecha.insert(0, datetime.datetime.now().strftime("%d/%m/%Y"))
        self.ent_fecha.grid(row=1, column=3, padx=5, pady=15, sticky="w")

        ctk.CTkLabel(form_card, text="Nivel de Gravedad:", text_color=self.COLOR_TEXT_MAIN, font=ctk.CTkFont(weight="bold")).grid(row=2, column=0, padx=(20, 5), pady=15, sticky="w")
        self.combo_gravedad = ctk.CTkComboBox(form_card, values=["Baja", "Media", "Alta", "Urgente"], width=300, height=36, state="readonly")
        self.combo_gravedad.grid(row=2, column=1, padx=5, pady=15, sticky="w")

        ctk.CTkLabel(form_card, text="Motivo / Descripción:", text_color=self.COLOR_TEXT_MAIN, font=ctk.CTkFont(weight="bold")).grid(row=3, column=0, padx=(20, 5), pady=15, sticky="nw")
        self.txt_motivo = ctk.CTkTextbox(form_card, width=300, height=130, border_width=1, border_color=self.COLOR_BORDER)
        self.txt_motivo.grid(row=3, column=1, padx=5, pady=15, sticky="w")

        ctk.CTkLabel(form_card, text="Observación / Sanción:", text_color=self.COLOR_TEXT_MAIN, font=ctk.CTkFont(weight="bold")).grid(row=3, column=2, padx=(20, 5), pady=15, sticky="nw")
        self.txt_sancion = ctk.CTkTextbox(form_card, width=300, height=130, border_width=1, border_color=self.COLOR_BORDER)
        self.txt_sancion.grid(row=3, column=3, padx=5, pady=15, sticky="w")

        btn_box = ctk.CTkFrame(form_card, fg_color="transparent")
        btn_box.grid(row=4, column=0, columnspan=4, pady=25)

        ctk.CTkButton(btn_box, text="💾 Emitir / Guardar Boleta", fg_color="#2563EB", hover_color="#1D4ED8", font=ctk.CTkFont(size=14, weight="bold"), width=220, height=40, command=self.guardar_boleta).pack(side="left", padx=10)
        ctk.CTkButton(btn_box, text="🧹 Nueva Boleta", fg_color="#64748B", hover_color="#475569", font=ctk.CTkFont(size=14, weight="bold"), width=150, height=40, command=self.limpiar_form_boleta).pack(side="left", padx=10)

        self.filtrar_combo_personal_boleta()

    def filtrar_combo_personal_boleta(self, event=None):
        q = self.ent_buscar_persona_boleta.get().lower().strip()
        lista = []
        for pid, pdata in self.bd["personal"].items():
            nombre = pdata.get('nombre', '')
            if not q or q in str(pid).lower() or q in nombre.lower():
                lista.append(f"{pid} - {nombre}")

        if lista:
            self.combo_personas.configure(values=lista)
            self.combo_personas.set(lista[0])
            self.al_seleccionar_persona_boleta(lista[0])
        else:
            self.combo_personas.configure(values=["No hay coincidencias"])
            self.combo_personas.set("No hay coincidencias")
            self.lbl_info_persona.configure(text="Cargo: - | Depto: -")
            self.persona_seleccionada_boleta = None

    def al_seleccionar_persona_boleta(self, choice):
        if " - " in choice:
            pid = choice.split(" - ")[0].strip()
            pdata = self.bd["personal"].get(pid, {})
            cargo = pdata.get("cargo", "N/A")
            depto = pdata.get("depto", "N/A")
            self.lbl_info_persona.configure(text=f"Cargo: {cargo} | Depto: {depto}")
            self.persona_seleccionada_boleta = pid

    def guardar_boleta(self):
        if not self.persona_seleccionada_boleta:
            messagebox.showwarning("Atención", "Debe seleccionar a una persona de la lista.")
            return

        pid = self.persona_seleccionada_boleta
        tipo = self.combo_tipo_boleta.get()
        fecha = self.ent_fecha.get().strip()
        gravedad = self.combo_gravedad.get()
        motivo = self.txt_motivo.get("1.0", "end").strip()
        sancion = self.txt_sancion.get("1.0", "end").strip()

        if not motivo:
            messagebox.showwarning("Atención", "El motivo no puede estar vacío.")
            return

        if self.boleta_id_actual is not None:
            for b in self.bd["boletas"]:
                if b["id_boleta"] == self.boleta_id_actual:
                    b.update({"persona_id": pid, "tipo": tipo, "fecha": fecha, "gravedad": gravedad, "motivo": motivo, "sancion": sancion})
                    break
            messagebox.showinfo("Éxito", f"Boleta N° {self.boleta_id_actual} actualizada.")
        else:
            nuevo_id = max([b["id_boleta"] for b in self.bd["boletas"]], default=0) + 1
            nueva = {"id_boleta": nuevo_id, "persona_id": pid, "tipo": tipo, "fecha": fecha, "gravedad": gravedad, "motivo": motivo, "sancion": sancion}
            self.bd["boletas"].append(nueva)
            messagebox.showinfo("Éxito", f"Boleta N° {nuevo_id} emitida.")

        guardar_bd(self.bd)
        self.limpiar_form_boleta()

    def limpiar_form_boleta(self):
        self.boleta_id_actual = None
        self.ent_buscar_persona_boleta.delete(0, "end")
        self.txt_motivo.delete("1.0", "end")
        self.txt_sancion.delete("1.0", "end")
        self.filtrar_combo_personal_boleta()

    # ==========================================
    # --- VISTA 4: HISTORIAL Y REPORTES ---
    # ==========================================
    def setup_vista_reportes(self):
        lbl_t = ctk.CTkLabel(self.vista_reportes, text="Historial General y Filtrado por Personal", font=ctk.CTkFont(size=24, weight="bold"), text_color=self.COLOR_TEXT_MAIN)
        lbl_t.pack(anchor="w", pady=(0, 15))

        toolbar = ctk.CTkFrame(self.vista_reportes, corner_radius=12, fg_color=self.COLOR_CARD, border_width=1, border_color=self.COLOR_BORDER)
        toolbar.pack(fill="x", pady=(0, 15))

        ctk.CTkLabel(toolbar, text="👤 Filtrar Usuario:", text_color=self.COLOR_TEXT_MAIN, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=(15, 5), pady=12)
        self.combo_filtro_usuario = ctk.CTkComboBox(toolbar, width=260, height=36, state="readonly", command=self.filtrar_boletas)
        self.combo_filtro_usuario.pack(side="left", padx=5, pady=12)

        ctk.CTkLabel(toolbar, text="🔍 Búsqueda Libre:", text_color=self.COLOR_TEXT_MAIN, font=ctk.CTkFont(weight="bold")).pack(side="left", padx=(15, 5), pady=12)
        self.ent_buscar = ctk.CTkEntry(toolbar, placeholder_text="Filtrar por texto o motivo...", width=220, height=36)
        self.ent_buscar.pack(side="left", padx=5, pady=12)
        self.ent_buscar.bind("<KeyRelease>", self.filtrar_boletas)

        ctk.CTkButton(toolbar, text="✏️ Editar", fg_color="#D97706", hover_color="#B45309", height=36, width=90, command=self.cargar_boleta_para_editar).pack(side="left", padx=5, pady=12)
        ctk.CTkButton(toolbar, text="🗑️ Eliminar", fg_color="#EF4444", hover_color="#DC2626", height=36, width=90, command=self.eliminar_boleta).pack(side="left", padx=5, pady=12)
        
        ctk.CTkButton(toolbar, text="📊 Exportar Excel", fg_color="#10B981", hover_color="#059669", font=ctk.CTkFont(weight="bold"), height=36, command=self.exportar_excel).pack(side="right", padx=15, pady=12)

        self.card_user_stats = ctk.CTkFrame(self.vista_reportes, corner_radius=12, fg_color="#EFF6FF", border_width=1, border_color="#BFDBFE")
        self.card_user_stats.pack(fill="x", pady=(0, 15))
        
        self.lbl_stats_usuario = ctk.CTkLabel(self.card_user_stats, text="Seleccione un usuario para ver su métrica individual.", font=ctk.CTkFont(size=13, weight="bold"), text_color="#1E40AF")
        self.lbl_stats_usuario.pack(padx=15, pady=10, anchor="w")

        table_card = ctk.CTkFrame(self.vista_reportes, corner_radius=12, fg_color=self.COLOR_CARD, border_width=1, border_color=self.COLOR_BORDER)
        table_card.pack(fill="both", expand=True)

        self.tv_boletas = ttk.Treeview(table_card, columns=("Nro", "Persona", "Tipo", "Fecha", "Gravedad", "Motivo", "Sancion"), show="headings")
        self.tv_boletas.heading("Nro", text="N°")
        self.tv_boletas.heading("Persona", text="Persona / ID")
        self.tv_boletas.heading("Tipo", text="Tipo")
        self.tv_boletas.heading("Fecha", text="Fecha")
        self.tv_boletas.heading("Gravedad", text="Gravedad")
        self.tv_boletas.heading("Motivo", text="Motivo")
        self.tv_boletas.heading("Sancion", text="Observación/Sanción")

        self.tv_boletas.column("Nro", width=60, anchor="center")
        self.tv_boletas.column("Persona", width=220)
        self.tv_boletas.column("Tipo", width=160)
        self.tv_boletas.column("Fecha", width=110, anchor="center")
        self.tv_boletas.column("Gravedad", width=120, anchor="center")
        self.tv_boletas.column("Motivo", width=250)
        self.tv_boletas.column("Sancion", width=250)

        self.tv_boletas.tag_configure("par", background="#F8FAFC")
        self.tv_boletas.tag_configure("impar", background="#FFFFFF")

        sb = ttk.Scrollbar(table_card, orient="vertical", command=self.tv_boletas.yview)
        self.tv_boletas.configure(yscrollcommand=sb.set)

        self.tv_boletas.pack(side="left", fill="both", expand=True, padx=(15, 0), pady=15)
        sb.pack(side="right", fill="y", padx=(0, 15), pady=15)

        self.tv_boletas.bind("<Double-1>", lambda event: self.cargar_boleta_para_editar())

    def actualizar_combo_filtro_usuarios(self):
        lista = ["-- Todos los Usuarios --"] + [f"{pid} - {pdata.get('nombre', '')}" for pid, pdata in self.bd["personal"].items()]
        self.combo_filtro_usuario.configure(values=lista)
        self.combo_filtro_usuario.set("-- Todos los Usuarios --")

    def actualizar_tabla_reportes(self, boletas_lista=None):
        for item in self.tv_boletas.get_children():
            self.tv_boletas.delete(item)

        datos = boletas_lista if boletas_lista is not None else self.bd["boletas"]

        for idx, b in enumerate(datos):
            pid = str(b.get("persona_id", ""))
            pnombre = self.bd["personal"].get(pid, {}).get("nombre", "Desconocido")
            tag = "par" if idx % 2 == 0 else "impar"
            
            grav = b.get("gravedad", "")
            iconos = {"Baja": "🟢 Baja", "Media": "🟡 Media", "Alta": "🟠 Alta", "Urgente": "🔴 Urgente"}
            grav_txt = iconos.get(grav, grav)

            self.tv_boletas.insert("", "end", values=(
                b.get("id_boleta"),
                f"{pnombre} ({pid})",
                b.get("tipo"),
                b.get("fecha"),
                grav_txt,
                b.get("motivo"),
                b.get("sancion", "")
            ), tags=(tag,))

    def filtrar_boletas(self, event=None):
        usr_sel = self.combo_filtro_usuario.get()
        query = self.ent_buscar.get().lower().strip()

        filtradas = self.bd["boletas"]

        if usr_sel != "-- Todos los Usuarios --" and " - " in usr_sel:
            pid_target = usr_sel.split(" - ")[0].strip()
            filtradas = [b for b in filtradas if str(b.get("persona_id")) == pid_target]
            
            tot_user = len(filtradas)
            criticos = sum(1 for b in filtradas if b.get("gravedad") in ["Alta", "Urgente"])
            self.lbl_stats_usuario.configure(text=f"📊 Resumen de {usr_sel}: Total Boletas: {tot_user} | Casos Críticos: {criticos}")
        else:
            self.lbl_stats_usuario.configure(text="Resumen General: Mostrando todas las boletas del sistema.")

        if query:
            res = []
            for b in filtradas:
                pid = str(b.get("persona_id", ""))
                pnombre = str(self.bd["personal"].get(pid, {}).get("nombre", "")).lower()
                tipo = str(b.get("tipo", "")).lower()
                motivo = str(b.get("motivo", "")).lower()

                if query in pid.lower() or query in pnombre or query in tipo or query in motivo:
                    res.append(b)
            filtradas = res

        self.actualizar_tabla_reportes(filtradas)

    def cargar_boleta_para_editar(self):
        sel = self.tv_boletas.selection()
        if not sel:
            messagebox.showwarning("Atención", "Seleccione una boleta de la lista.")
            return

        nro_boleta = self.tv_boletas.item(sel[0])["values"][0]
        boleta_encontrada = next((b for b in self.bd["boletas"] if b["id_boleta"] == nro_boleta), None)

        if boleta_encontrada:
            self.boleta_id_actual = boleta_encontrada["id_boleta"]
            pid = str(boleta_encontrada["persona_id"])

            self.ent_buscar_persona_boleta.delete(0, "end")
            self.ent_buscar_persona_boleta.insert(0, pid)
            self.filtrar_combo_personal_boleta()

            self.combo_tipo_boleta.set(boleta_encontrada.get("tipo", ""))
            self.ent_fecha.delete(0, "end")
            self.ent_fecha.insert(0, boleta_encontrada.get("fecha", ""))
            self.combo_gravedad.set(boleta_encontrada.get("gravedad", ""))

            self.txt_motivo.delete("1.0", "end")
            self.txt_motivo.insert("1.0", boleta_encontrada.get("motivo", ""))

            self.txt_sancion.delete("1.0", "end")
            self.txt_sancion.insert("1.0", boleta_encontrada.get("sancion", ""))

            self.mostrar_vista("boletas")

    def eliminar_boleta(self):
        sel = self.tv_boletas.selection()
        if not sel:
            messagebox.showwarning("Atención", "Seleccione una boleta.")
            return

        nro_boleta = self.tv_boletas.item(sel[0])["values"][0]
        if messagebox.askyesno("Confirmar", f"¿Eliminar la Boleta N° {nro_boleta}?"):
            self.bd["boletas"] = [b for b in self.bd["boletas"] if b["id_boleta"] != nro_boleta]
            guardar_bd(self.bd)
            self.actualizar_tabla_reportes()

    # ==========================================
    # --- EXPORTACIÓN EXCEL ---
    # ==========================================
    def exportar_excel(self):
        if not self.bd["boletas"]:
            messagebox.showwarning("Atención", "No hay registros de boletas para exportar.")
            return

        nombre_sugerido = f"Reporte_Disciplinario_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        ruta_archivo = filedialog.asksaveasfilename(
            title="Guardar Reporte Excel",
            initialfile=nombre_sugerido,
            defaultextension=".xlsx",
            filetypes=[("Archivos de Excel", "*.xlsx")]
        )

        if not ruta_archivo:
            return

        filas = []
        for b in self.bd["boletas"]:
            pid = str(b.get("persona_id", ""))
            pdata = self.bd["personal"].get(pid, {})
            
            filas.append({
                "N° Boleta": b.get("id_boleta"),
                "ID Persona": pid,
                "Nombre Persona": pdata.get("nombre", "Desconocido"),
                "Cargo/Curso": pdata.get("cargo", ""),
                "Departamento": pdata.get("depto", ""),
                "Tipo de Boleta": b.get("tipo", ""),
                "Fecha": b.get("fecha", ""),
                "Nivel Gravedad": b.get("gravedad", ""),
                "Motivo / Descripción": b.get("motivo", ""),
                "Observación / Sanción": b.get("sancion", "")
            })

        df = pd.DataFrame(filas)

        try:
            with pd.ExcelWriter(ruta_archivo, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name="Historial", startrow=3)
                
                worksheet = writer.sheets["Historial"]
                
                COLOR_AZUL_HEADER = "0F172A"
                COLOR_PAR = "F8FAFC"
                COLOR_IMPAR = "FFFFFF"
                COLOR_BORDER = "CBD5E1"
                
                font_titulo = Font(name="Segoe UI", size=16, bold=True, color="0F172A")
                font_subtitulo = Font(name="Segoe UI", size=10, italic=True, color="64748B")
                font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                font_datos = Font(name="Segoe UI", size=10, color="0F172A")
                
                fill_header = PatternFill(start_color=COLOR_AZUL_HEADER, end_color=COLOR_AZUL_HEADER, fill_type="solid")
                fill_par = PatternFill(start_color=COLOR_PAR, end_color=COLOR_PAR, fill_type="solid")
                fill_impar = PatternFill(start_color=COLOR_IMPAR, end_color=COLOR_IMPAR, fill_type="solid")
                
                fills_gravedad = {
                    "Baja": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
                    "Media": PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),
                    "Alta": PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
                    "Urgente": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                }
                fonts_gravedad = {
                    "Baja": Font(name="Segoe UI", size=10, bold=True, color="166534"),
                    "Media": Font(name="Segoe UI", size=10, bold=True, color="854D0E"),
                    "Alta": Font(name="Segoe UI", size=10, bold=True, color="9A3412"),
                    "Urgente": Font(name="Segoe UI", size=10, bold=True, color="991B1B")
                }

                border_thin = Border(
                    left=Side(style='thin', color=COLOR_BORDER),
                    right=Side(style='thin', color=COLOR_BORDER),
                    top=Side(style='thin', color=COLOR_BORDER),
                    bottom=Side(style='thin', color=COLOR_BORDER)
                )

                worksheet["A1"] = "REPORTE GENERAL DE CONTROL DISCIPLINARIO"
                worksheet["A1"].font = font_titulo
                
                worksheet["A2"] = f"Generado el: {datetime.datetime.now().strftime('%d/%m/%Y a las %H:%M hrs')}"
                worksheet["A2"].font = font_subtitulo

                header_row = 4
                max_row = worksheet.max_row
                max_col = worksheet.max_column

                for col in range(1, max_col + 1):
                    cell = worksheet.cell(row=header_row, column=col)
                    cell.font = font_header
                    cell.fill = fill_header
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = border_thin

                for row in range(header_row + 1, max_row + 1):
                    fill_row = fill_par if row % 2 == 0 else fill_impar
                    for col in range(1, max_col + 1):
                        cell = worksheet.cell(row=row, column=col)
                        cell.font = font_datos
                        cell.fill = fill_row
                        cell.border = border_thin
                        
                        if col in [1, 2, 7]:
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                        elif col == 8:
                            val = str(cell.value)
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                            if val in fills_gravedad:
                                cell.fill = fills_gravedad[val]
                                cell.font = fonts_gravedad[val]
                        else:
                            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

                worksheet.row_dimensions[header_row].height = 28
                for row in range(header_row + 1, max_row + 1):
                    worksheet.row_dimensions[row].height = 22

                for col in worksheet.columns:
                    max_len = 0
                    col_letter = get_column_letter(col[0].column)
                    for cell in col:
                        if cell.row < header_row:
                            continue
                        if cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    ancho_calculado = max(max_len + 4, 12)
                    worksheet.column_dimensions[col_letter].width = min(ancho_calculado, 35)

            messagebox.showinfo("Éxito", f"Reporte guardado correctamente en:\n{ruta_archivo}")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo guardar el archivo Excel.\nDetalle: {str(e)}")

if __name__ == "__main__":
    app = DashboardAdmin()
    app.mainloop()